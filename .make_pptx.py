#!/usr/bin/env python3
"""Never Miss Bus — presentation deck (PPTX)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import os

PRIMARY = RGBColor(0x25, 0x57, 0xD6)
PRIMARY_DARK = RGBColor(0x1A, 0x3F, 0xA0)
PRIMARY_SOFT = RGBColor(0xE8, 0xEE, 0xFC)
ACCENT = RGBColor(0xFF, 0xB3, 0x00)
ACCENT_DARK = RGBColor(0x8A, 0x5B, 0x00)
SUCCESS = RGBColor(0x1E, 0x8E, 0x3E)
DANGER = RGBColor(0xC5, 0x22, 0x1F)
TEXT = RGBColor(0x17, 0x23, 0x3B)
TEXT2 = RGBColor(0x5A, 0x67, 0x79)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF6, 0xF8, 0xFC)
DIVIDER = RGBColor(0xE4, 0xE9, 0xF2)

MOCK = "/home/user/never_miss_bus/docs/mockups"
OUT = "/home/user/never_miss_bus/docs/Never_Miss_Bus_Presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]

def slide(bg=BG):
    s = prs.slides.add_slide(blank)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s

def box(s, x, y, w, h, fill=None, line=None, radius=True):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    if radius:
        try:
            shp.adjustments[0] = 0.08
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1.2)
    shp.shadow.inherit = False
    return shp

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        for t, size, bold, color in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = "Calibri"
    return tb

def title_bar(s, kicker, title):
    box(s, 0, 0, SW, Inches(0.42), fill=PRIMARY, radius=False)
    text(s, Inches(0.35), Inches(0.04), Inches(9), Inches(0.35),
         [[(kicker, 12, True, WHITE)]])
    text(s, Inches(0.6), Inches(0.62), Inches(12), Inches(0.9),
         [[(title, 32, True, PRIMARY_DARK)]])

def bullets(s, x, y, w, h, items, size=16, gap=10):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for head, body in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = "•  " + head
        r.font.size = Pt(size); r.font.bold = True
        r.font.color.rgb = TEXT; r.font.name = "Calibri"
        if body:
            r2 = p.add_run(); r2.text = " — " + body
            r2.font.size = Pt(size); r2.font.bold = False
            r2.font.color.rgb = TEXT2; r2.font.name = "Calibri"
    return tb

def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt

# ═══ SLIDE 1 — Title ═══
s = slide(PRIMARY)
box(s, 0, SH - Inches(1.1), SW, Inches(1.1), fill=PRIMARY_DARK, radius=False)
c = s.shapes.add_shape(MSO_SHAPE.OVAL, SW/2 - Inches(0.9),
                       Inches(0.75), Inches(1.8), Inches(1.8))
c.fill.solid(); c.fill.fore_color.rgb = WHITE; c.line.fill.background()
c.shadow.inherit = False
b = box(s, SW/2 - Inches(0.62), Inches(1.28), Inches(1.24), Inches(0.62),
        fill=ACCENT)
for i in range(3):
    w = box(s, SW/2 - Inches(0.47) + Inches(0.34)*i, Inches(1.38),
            Inches(0.24), Inches(0.18), fill=RGBColor(0xE3, 0xF2, 0xFD))
for dx in (-0.35, 0.35):
    wh = s.shapes.add_shape(MSO_SHAPE.OVAL, SW/2 + Inches(dx) - Inches(0.09),
                            Inches(1.82), Inches(0.18), Inches(0.18))
    wh.fill.solid(); wh.fill.fore_color.rgb = RGBColor(0x37, 0x47, 0x4F)
    wh.line.fill.background(); wh.shadow.inherit = False
text(s, 0, Inches(2.85), SW, Inches(1),
     [[("NEVER MISS BUS", 54, True, WHITE)]], align=PP_ALIGN.CENTER)
text(s, 0, Inches(3.85), SW, Inches(0.6),
     [[("Secure School Bus Live-Tracking App", 24, False,
        RGBColor(0xCF, 0xE0, 0xFF))]], align=PP_ALIGN.CENTER)
bdg = box(s, SW/2 - Inches(2.2), Inches(4.6), Inches(4.4), Inches(0.55),
          fill=ACCENT)
text(s, SW/2 - Inches(2.2), Inches(4.66), Inches(4.4), Inches(0.45),
     [[("SRBS INTERNATIONAL SCHOOL", 16, True, RGBColor(0x5B, 0x3C, 0x00))]],
     align=PP_ALIGN.CENTER)
text(s, 0, Inches(5.5), SW, Inches(0.5),
     [[("Android + iOS  ·  Flutter + Firebase  ·  Student · Driver · Admin",
        16, False, RGBColor(0xCF, 0xE0, 0xFF))]], align=PP_ALIGN.CENTER)
text(s, 0, SH - Inches(0.85), SW, Inches(0.5),
     [[("Presented by:  [Apna naam yahan]     |     August 2026", 14, False,
        WHITE)]], align=PP_ALIGN.CENTER)
notes(s, "OPENING (30 sec): 'Good morning! Kya aapke bachhe ne kabhi bus miss ki hai? "
         "Ya aap sardi/baarish mein stop pe khade intezaar karte rahe ho bina jaane bus kahan hai? "
         "Main aapko dikhata hoon Never Miss Bus — hamare school ke liye ek secure live bus tracking app.'")

# ═══ SLIDE 2 — Problem ═══
s = slide()
title_bar(s, "THE PROBLEM", "Roz ki pareshani — bus ka intezaar")
probs = [
    ("Students stop pe intezaar karte hain", "pata nahi bus 2 minute door hai ya 20"),
    ("Parents ko tension rehti hai", "bachha nikla, bus aayi ya nahi — koi jaankari nahi"),
    ("Bus miss = poora din kharab", "late arrival, extra trips, phone calls"),
    ("School office pe calls ka bojh", "'Bus kahan hai?' — har roz wahi sawaal"),
    ("Generic GPS apps safe nahi hain", "public tracking links = privacy risk for children"),
]
y = Inches(1.7)
for head, body in probs:
    box(s, Inches(0.6), y, Inches(12.1), Inches(0.92), fill=WHITE, line=DIVIDER)
    text(s, Inches(0.95), y + Inches(0.12), Inches(11.5), Inches(0.7),
         [[(head, 18, True, TEXT)], [(body, 14, False, TEXT2)]], space_after=2)
    y += Inches(1.06)
notes(s, "PROBLEM (1 min): Relatable banao. Audience se poochho — 'Kitne logon ne yeh feel kiya hai?' "
         "Last point important hai: normal GPS-sharing apps bachhon ke liye unsafe hain kyunki koi bhi link se location dekh sakta hai.")

# ═══ SLIDE 3 — Solution ═══
s = slide()
title_bar(s, "THE SOLUTION", "Never Miss Bus — ek app, teen roles")
cards = [
    ("STUDENT / PARENT", PRIMARY,
     ["Apni bus LIVE map pe dekho", "Route, stops, apna stop", "'~8 min mein aa rahi hai' ETA",
      "Bus approaching alerts"]),
    ("DRIVER", SUCCESS,
     ["Ek button: START TRIP", "GPS automatic share hota hai", "Driving ke waqt kuch nahi karna",
      "END TRIP — tracking band"]),
    ("SCHOOL ADMIN", ACCENT_DARK,
     ["Buses, routes, stops manage", "Students/drivers assign karo", "Saari buses ek map pe",
      "Announcements bhejo"]),
]
x = Inches(0.6)
for title_, col, items in cards:
    box(s, x, Inches(1.75), Inches(4.0), Inches(4.9), fill=WHITE, line=DIVIDER)
    box(s, x, Inches(1.75), Inches(4.0), Inches(0.7), fill=col)
    text(s, x, Inches(1.85), Inches(4.0), Inches(0.5),
         [[(title_, 17, True, WHITE)]], align=PP_ALIGN.CENTER)
    ty = Inches(2.7)
    for it in items:
        text(s, x + Inches(0.3), ty, Inches(3.5), Inches(0.8),
             [[("✓  ", 15, True, col), (it, 15, False, TEXT)]])
        ty += Inches(0.78)
    x += Inches(4.25)
text(s, Inches(0.6), Inches(6.85), Inches(12), Inches(0.5),
     [[("Har role ko sirf wahi dikhta hai jo uske liye hai — backend rules se enforced.",
        15, True, PRIMARY_DARK)]], align=PP_ALIGN.CENTER)
notes(s, "SOLUTION (1 min): Teen roles clearly batao. Key line: 'Student sirf APNI bus dekh sakta hai — "
         "yeh app ka rule nahi, server ka rule hai. Koi hack karke bhi doosri bus nahi dekh sakta.'")

# ═══ SLIDE 4 — Student app (mockups) ═══
s = slide()
title_bar(s, "STUDENT EXPERIENCE", "Bachhon ke liye — simple aur friendly")
s.shapes.add_picture(f"{MOCK}/student_home.png", Inches(1.1), Inches(1.55),
                     height=Inches(5.7))
s.shapes.add_picture(f"{MOCK}/student_map.png", Inches(4.4), Inches(1.55),
                     height=Inches(5.7))
bullets(s, Inches(7.9), Inches(1.9), Inches(5.0), Inches(5), [
    ("Greeting + apni bus card", "bus number, plate, ON TRIP status"),
    ("Bada 'Track My Bus' button", "ek tap — live map khul jaata hai"),
    ("School-bus marker 🚌", "route line, saare stops, apna stop star ke saath"),
    ("Honest ETA", "'approximately 8 min' — kabhi jhootha vaada nahi"),
    ("Alerts", "trip started, bus approaching, bus reached stop"),
    ("Agar signal chala jaye", "'Last updated 2 min ago' dikhata hai — purani location ko live nahi bolta"),
], size=15, gap=12)
notes(s, "STUDENT (1.5 min): Mockups pe point karo. Honest-status wali baat highlight karo — "
         "'Agar driver ka network chala jaye, app jhooth nahi bolti. Saaf batati hai location kitni purani hai.' "
         "Yeh trust build karta hai.")

# ═══ SLIDE 5 — Driver + Admin (mockups) ═══
s = slide()
title_bar(s, "DRIVER & ADMIN", "Driver ke liye 1 button · Admin ke liye full control")
s.shapes.add_picture(f"{MOCK}/driver_home.png", Inches(1.0), Inches(1.55),
                     height=Inches(5.7))
s.shapes.add_picture(f"{MOCK}/admin_dashboard.png", Inches(4.3), Inches(1.55),
                     height=Inches(5.7))
bullets(s, Inches(7.8), Inches(1.9), Inches(5.1), Inches(5), [
    ("Driver: START TRIP dabao, bas", "GPS apne aap share hota hai — screen band ho toh bhi"),
    ("Driving ke waqt zero distraction", "koi interaction ki zaroorat nahi — safety first"),
    ("Admin: poora transport structure", "Bus → Driver → Route → Stops → Students"),
    ("Live monitoring", "saari active buses ek hi map pe"),
    ("Ek click mein account disable", "koi phone kho jaye toh turant access band"),
    ("Activity logs", "kaun kya kab kiya — sab recorded"),
], size=15, gap=12)
notes(s, "DRIVER+ADMIN (1.5 min): Driver safety pe zor do — 'Driver ko driving ke time phone chhoona hi nahi padta.' "
         "Admin side pe school ka control dikhao — 'School ke paas poora control hai, parents ko sirf dekhne ka access.'")

# ═══ SLIDE 6 — Security ═══
s = slide()
title_bar(s, "SECURITY & PRIVACY", "Bachhon ki safety = sabse pehli priority")
rows = [
    ("Student sirf apni bus dekh sakta hai", "server-side rules — app hack karke bhi bypass nahi hota"),
    ("Koi public tracking link nahi", "bina school-account login ke kuch nahi dikhta"),
    ("Accounts sirf school banata hai", "koi khud register nahi kar sakta, koi apni bus change nahi kar sakta"),
    ("Students ka personal data minimal", "sirf naam, class, bus/stop — phone number tak store nahi hota"),
    ("Location history store nahi hoti", "trip khatam = location data delete"),
    ("Passwords encrypted (Google Firebase)", "kabhi plain text mein store nahi hote"),
    ("Har admin action ka record", "tamper-proof audit log — koi mita nahi sakta"),
]
y = Inches(1.65)
for head, body in rows:
    box(s, Inches(0.6), y, Inches(12.1), Inches(0.68), fill=WHITE, line=DIVIDER)
    text(s, Inches(0.8), y + Inches(0.07), Inches(0.5), Inches(0.5),
         [[("🔒", 16, False, TEXT)]])
    text(s, Inches(1.35), y + Inches(0.05), Inches(5.4), Inches(0.55),
         [[(head, 14.5, True, TEXT)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(6.9), y + Inches(0.05), Inches(5.7), Inches(0.55),
         [[(body, 13, False, TEXT2)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.78)
notes(s, "SECURITY (1.5 min): Yeh slide parents/principal ke liye sabse important hai. "
         "Bolo: 'Yeh WhatsApp location sharing ya generic tracker se bilkul alag hai. "
         "Yahan security database ke level pe hai — Google ke Firebase rules — "
         "app todne se bhi data nahi milta.'")

# ═══ SLIDE 7 — Technology ═══
s = slide()
title_bar(s, "TECHNOLOGY", "Modern, proven, scalable stack")
tech = [
    ("Flutter", "Ek codebase → Android + iOS dono. Google ki technology.", PRIMARY),
    ("Firebase", "Google ka secure backend — Auth, database, notifications.", ACCENT_DARK),
    ("Google Maps", "Wahi maps jo sab roz use karte hain.", SUCCESS),
    ("Cloud Functions", "Server-side logic — accounts, alerts, security.", PRIMARY_DARK),
]
x = Inches(0.6)
for name, desc, col in tech:
    box(s, x, Inches(1.75), Inches(2.95), Inches(2.1), fill=WHITE, line=DIVIDER)
    box(s, x, Inches(1.75), Inches(2.95), Inches(0.14), fill=col)
    text(s, x + Inches(0.25), Inches(2.1), Inches(2.5), Inches(0.5),
         [[(name, 19, True, TEXT)]])
    text(s, x + Inches(0.25), Inches(2.65), Inches(2.5), Inches(1.1),
         [[(desc, 13, False, TEXT2)]])
    x += Inches(3.13)
box(s, Inches(0.6), Inches(4.25), Inches(12.1), Inches(2.5), fill=PRIMARY_SOFT)
text(s, Inches(1.0), Inches(4.5), Inches(11.4), Inches(2.1), [
    [("Quality metrics:", 17, True, PRIMARY_DARK)],
    [("✓  63 Dart files, ~8,000 lines — clean architecture, 0 analyzer issues", 15, False, TEXT)],
    [("✓  15/15 automated tests passing (freshness honesty, ETA guardrails, security parsing)", 15, False, TEXT)],
    [("✓  Working APK built & verified — aaj hi install karke dikha sakta hoon", 15, False, TEXT)],
    [("✓  Running cost: chhote school ke liye lagbhag ₹0/month (Firebase free limits ke andar)", 15, False, TEXT)],
], space_after=8)
notes(s, "TECH (1 min): Non-technical audience ke liye simple rakho — 'Google ki hi technologies pe bana hai. "
         "Cost ki baat zaroor karo: Firebase ka free tier chhote school ke liye kaafi hai.'")

# ═══ SLIDE 8 — Live Demo ═══
s = slide()
title_bar(s, "LIVE DEMO", "Ab main aapko chala ke dikhata hoon")
steps = [
    ("1", "Admin login", "Bus 3 banayi, route + 4 stops add kiye"),
    ("2", "Student assign", "Aarav ko Bus 3 + Rajpur Road stop diya"),
    ("3", "Driver phone", "START TRIP dabaya — GPS on"),
    ("4", "Student phone", "Track My Bus → bus LIVE chal rahi hai 🚌"),
    ("5", "Notification", "'Bus approaching' alert aaya"),
    ("6", "END TRIP", "tracking band — student ko saaf dikh gaya"),
]
y = Inches(1.7)
for num, head, body in steps:
    box(s, Inches(1.2), y, Inches(10.9), Inches(0.78), fill=WHITE, line=DIVIDER)
    ovl = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.45), y + Inches(0.14),
                             Inches(0.5), Inches(0.5))
    ovl.fill.solid(); ovl.fill.fore_color.rgb = PRIMARY
    ovl.line.fill.background(); ovl.shadow.inherit = False
    tf = ovl.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run(); r.text = num
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
    text(s, Inches(2.2), y + Inches(0.09), Inches(3.2), Inches(0.6),
         [[(head, 16, True, TEXT)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(5.5), y + Inches(0.09), Inches(6.3), Inches(0.6),
         [[(body, 14, False, TEXT2)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.88)
notes(s, "DEMO (3-4 min): Do phone use karo — ek driver, ek student. Agar live demo risky lage "
         "toh pehle se screen-recording bana lo (backup video). Demo se pehle sab login karke ready rakho. "
         "Wow moment: student phone pe bus marker MOVE hota hua dikhao.")

# ═══ SLIDE 9 — Roadmap / Closing ═══
s = slide()
title_bar(s, "NEXT STEPS", "Launch plan aur aage ka roadmap")
cols = [
    ("Phase 1 — Pilot (2 hafte)", PRIMARY,
     ["1 bus, 1 driver, ~20 students", "Feedback lo, tuning karo"]),
    ("Phase 2 — Full rollout", SUCCESS,
     ["Saari buses + routes", "Parents ko onboarding SMS/note"]),
    ("Phase 3 — Future ideas", ACCENT_DARK,
     ["Parent-specific accounts", "Attendance on boarding", "Hindi language option"]),
]
x = Inches(0.6)
for title_, col, items in cols:
    box(s, x, Inches(1.7), Inches(4.0), Inches(3.1), fill=WHITE, line=DIVIDER)
    box(s, x, Inches(1.7), Inches(4.0), Inches(0.6), fill=col)
    text(s, x + Inches(0.2), Inches(1.78), Inches(3.6), Inches(0.45),
         [[(title_, 15, True, WHITE)]])
    ty = Inches(2.55)
    for it in items:
        text(s, x + Inches(0.3), ty, Inches(3.5), Inches(0.6),
             [[("•  " + it, 14, False, TEXT)]])
        ty += Inches(0.55)
    x += Inches(4.25)
box(s, Inches(0.6), Inches(5.15), Inches(12.1), Inches(1.6), fill=PRIMARY)
text(s, Inches(0.6), Inches(5.45), Inches(12.1), Inches(1.2), [
    [("Never Miss Bus — kyunki har bachhe ka time, safety aur sukoon important hai.",
      20, True, WHITE)],
    [("Questions?  Main demo ke liye taiyaar hoon. 🚌", 16, False,
      RGBColor(0xCF, 0xE0, 0xFF))],
], align=PP_ALIGN.CENTER, space_after=10)
notes(s, "CLOSING (30 sec): Confident khatam karo. Pilot ka idea do — 'Hum ek bus se shuru kar sakte hain, "
         "2 hafte mein results dikha denge.' Phir Q&A kholo. Common questions ki taiyaari DEMO_SCRIPT mein hai.")

prs.save(OUT)
print("saved:", OUT, os.path.getsize(OUT)//1024, "KB")
