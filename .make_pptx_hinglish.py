#!/usr/bin/env python3
"""Never Miss Bus — Hinglish (Hindi-English mixed) presentation deck."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

PRIMARY = RGBColor(0x25, 0x57, 0xD6)
PRIMARY_DARK = RGBColor(0x1A, 0x3F, 0xA0)
PRIMARY_SOFT = RGBColor(0xE8, 0xEE, 0xFC)
ACCENT = RGBColor(0xFF, 0xB3, 0x00)
ACCENT_DARK = RGBColor(0x8A, 0x5B, 0x00)
SUCCESS = RGBColor(0x1E, 0x8E, 0x3E)
SUCCESS_SOFT = RGBColor(0xE2, 0xF3, 0xE7)
DANGER = RGBColor(0xC5, 0x22, 0x1F)
TEXT = RGBColor(0x17, 0x23, 0x3B)
TEXT2 = RGBColor(0x5A, 0x67, 0x79)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF6, 0xF8, 0xFC)
DIVIDER = RGBColor(0xE4, 0xE9, 0xF2)
PURPLE = RGBColor(0x7B, 0x61, 0xFF)

MOCK = "/home/user/never_miss_bus/docs/mockups"
OUT = "/home/user/never_miss_bus/docs/Never_Miss_Bus_Presentation_Hinglish.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]

def slide(bg=BG):
    s = prs.slides.add_slide(blank)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background(); r.shadow.inherit = False
    return s

def box(s, x, y, w, h, fill=None, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try: shp.adjustments[0] = 0.08
    except Exception: pass
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(1.2)
    shp.shadow.inherit = False
    return shp

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align; p.space_after = Pt(space_after)
        for t, size, bold, color in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = "Calibri"
    return tb

def title_bar(s, kicker, title):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.42))
    bar.fill.solid(); bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background(); bar.shadow.inherit = False
    text(s, Inches(0.35), Inches(0.04), Inches(9), Inches(0.35),
         [[(kicker, 12, True, WHITE)]])
    text(s, Inches(0.6), Inches(0.62), Inches(12.2), Inches(0.9),
         [[(title, 30, True, PRIMARY_DARK)]])

def bullets(s, x, y, w, h, items, size=15, gap=10):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for head, body in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        r = p.add_run(); r.text = "•  " + head
        r.font.size = Pt(size); r.font.bold = True
        r.font.color.rgb = TEXT; r.font.name = "Calibri"
        if body:
            r2 = p.add_run(); r2.text = " — " + body
            r2.font.size = Pt(size); r2.font.bold = False
            r2.font.color.rgb = TEXT2; r2.font.name = "Calibri"
    return tb

def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt

# ═══ SLIDE 1 — Title ═══
s = slide(PRIMARY)
b2 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SH - Inches(1.1), SW, Inches(1.1))
b2.fill.solid(); b2.fill.fore_color.rgb = PRIMARY_DARK
b2.line.fill.background(); b2.shadow.inherit = False
c = s.shapes.add_shape(MSO_SHAPE.OVAL, SW/2 - Inches(0.9), Inches(0.7),
                       Inches(1.8), Inches(1.8))
c.fill.solid(); c.fill.fore_color.rgb = WHITE
c.line.fill.background(); c.shadow.inherit = False
bus = box(s, SW/2 - Inches(0.62), Inches(1.22), Inches(1.24), Inches(0.62),
          fill=ACCENT)
for i in range(3):
    box(s, SW/2 - Inches(0.47) + Inches(0.34)*i, Inches(1.32),
        Inches(0.24), Inches(0.18), fill=RGBColor(0xE3, 0xF2, 0xFD))
for dx in (-0.35, 0.35):
    wh = s.shapes.add_shape(MSO_SHAPE.OVAL, SW/2 + Inches(dx) - Inches(0.09),
                            Inches(1.76), Inches(0.18), Inches(0.18))
    wh.fill.solid(); wh.fill.fore_color.rgb = RGBColor(0x37, 0x47, 0x4F)
    wh.line.fill.background(); wh.shadow.inherit = False
text(s, 0, Inches(2.75), SW, Inches(1),
     [[("NEVER MISS BUS", 52, True, WHITE)]], align=PP_ALIGN.CENTER)
text(s, 0, Inches(3.7), SW, Inches(0.6),
     [[("School Bus Live-Tracking App — Ab Bus Miss Nahi Hogi! 🚌", 22,
        False, RGBColor(0xCF, 0xE0, 0xFF))]], align=PP_ALIGN.CENTER)
bdg = box(s, SW/2 - Inches(2.3), Inches(4.45), Inches(4.6), Inches(0.55),
          fill=ACCENT)
text(s, SW/2 - Inches(2.3), Inches(4.51), Inches(4.6), Inches(0.45),
     [[("SRBS INTERNATIONAL SCHOOL", 16, True, RGBColor(0x5B, 0x3C, 0x00))]],
     align=PP_ALIGN.CENTER)
text(s, 0, Inches(5.35), SW, Inches(0.5),
     [[("📱 Android + iOS (iPhone) dono ke liye  ·  Ek hi app — "
        "Student, Driver, Admin", 16, False, RGBColor(0xCF, 0xE0, 0xFF))]],
     align=PP_ALIGN.CENTER)
text(s, 0, SH - Inches(0.85), SW, Inches(0.5),
     [[("Presented by:  [Apna naam]     |     Version 1.2.0     |     2026",
        14, False, WHITE)]], align=PP_ALIGN.CENTER)
notes(s, "OPENING (30 sec): 'Good morning! Ek sawaal — kitne logon ke ghar mein roz subah "
         "yeh tension hoti hai: bus aayi ya nikal gayi? Aaj main aapko dikhata hoon iska "
         "solution — Never Miss Bus, hamare school ka apna secure bus tracking app, jo "
         "Android aur iPhone dono par chalta hai.'")

# ═══ SLIDE 2 — Problem ═══
s = slide()
title_bar(s, "THE PROBLEM / SAMASYA", "Roz Subah Ki Tension — Bus Kahan Hai?")
probs = [
    ("Students stop par wait karte hain", "pata nahi bus 2 minute door hai ya 20 minute"),
    ("Parents pareshan rehte hain", "bachcha nikla, bus aayi ya nahi — koi information nahi"),
    ("Bus miss = poora din kharab", "late arrival, extra trip, phone calls"),
    ("School office par calls ka load", "'Bus kahan hai?' — daily wahi question"),
    ("Normal GPS apps children ke liye unsafe hain", "public tracking links = privacy risk"),
]
y = Inches(1.7)
for head, body in probs:
    box(s, Inches(0.6), y, Inches(12.1), Inches(0.92), fill=WHITE, line=DIVIDER)
    text(s, Inches(0.95), y + Inches(0.12), Inches(11.5), Inches(0.7),
         [[(head, 18, True, TEXT)], [(body, 14, False, TEXT2)]], space_after=2)
    y += Inches(1.06)
notes(s, "PROBLEM (1 min): Relatable banao — audience se poochho 'Kitne logon ne yeh "
         "feel kiya hai?' Last point important: normal GPS-sharing apps bachchon ke liye "
         "unsafe hain kyunki koi bhi link se location dekh sakta hai.")

# ═══ SLIDE 3 — Solution: 3 roles ═══
s = slide()
title_bar(s, "THE SOLUTION / SAMADHAN", "Ek App — Teen Roles, Sabka Kaam Easy")
cards = [
    ("STUDENT / PARENT", PRIMARY,
     ["Apni bus LIVE map par dekho", "Class + Roll number se login — email nahi chahiye!",
      "Route, stops, apna stop ⭐", "'~8 min mein aa rahi hai' — ETA",
      "Bus approaching alerts"]),
    ("DRIVER", SUCCESS,
     ["Sirf ek button: START TRIP", "GPS automatic share hota hai",
      "Driving ke time kuch touch nahi karna", "Screen band ho toh bhi chalta hai",
      "END TRIP — tracking band"]),
    ("SCHOOL ADMIN", ACCENT_DARK,
     ["Buses, routes, stops manage karo", "Students class-wise add/edit karo",
      "Saari buses ek map par LIVE", "Kaun logged in hai — sab dikhega",
      "Force logout + password reset"]),
]
x = Inches(0.6)
for title_, col, items in cards:
    box(s, x, Inches(1.7), Inches(4.0), Inches(5.0), fill=WHITE, line=DIVIDER)
    hdr = box(s, x, Inches(1.7), Inches(4.0), Inches(0.65), fill=col)
    text(s, x, Inches(1.79), Inches(4.0), Inches(0.5),
         [[(title_, 16, True, WHITE)]], align=PP_ALIGN.CENTER)
    ty = Inches(2.55)
    for it in items:
        text(s, x + Inches(0.28), ty, Inches(3.55), Inches(0.75),
             [[("✓  ", 13.5, True, col), (it, 13.5, False, TEXT)]])
        ty += Inches(0.8)
    x += Inches(4.25)
notes(s, "SOLUTION (1.5 min): Teen roles clearly batao. Key line: 'Student ko email yaad "
         "rakhne ki zaroorat nahi — bas Class, Section, Roll number aur password. Aur "
         "student sirf APNI bus dekh sakta hai — yeh rule server par hai, hack karke bhi "
         "nahi tuteta.'")

# ═══ SLIDE 4 — Student app mockups ═══
s = slide()
title_bar(s, "STUDENT EXPERIENCE", "Bachchon Ke Liye — Simple, Friendly, Safe")
if os.path.exists(f"{MOCK}/student_home.png"):
    s.shapes.add_picture(f"{MOCK}/student_home.png", Inches(1.0), Inches(1.55),
                         height=Inches(5.6))
if os.path.exists(f"{MOCK}/student_map.png"):
    s.shapes.add_picture(f"{MOCK}/student_map.png", Inches(4.3), Inches(1.55),
                         height=Inches(5.6))
bullets(s, Inches(7.8), Inches(1.8), Inches(5.1), Inches(5.3), [
    ("Login super easy", "Class dropdown + Section + Roll number + password — bas!"),
    ("Home screen", "greeting, apni bus card, ON TRIP status, bada 'Track My Bus' button"),
    ("Live Map 🚌", "school-bus marker, route line, saare stops, apna stop star ke saath"),
    ("Honest ETA", "'approximately 8 min' — app kabhi jhootha promise nahi karti"),
    ("Signal chala jaye toh?", "'Last updated 2 min ago' dikhata hai — purani location ko LIVE nahi bolta"),
    ("Alerts", "trip started, bus approaching, bus reached stop — automatic"),
], size=14, gap=11)
notes(s, "STUDENT (1.5 min): Mockups par point karo. Honest-status highlight karo: "
         "'Agar driver ka network chala jaye, app clearly batati hai location kitni "
         "purani hai. Yeh trust build karta hai.'")

# ═══ SLIDE 5 — Driver + Admin mockups ═══
s = slide()
title_bar(s, "DRIVER & ADMIN", "Driver: 1 Button · Admin: Full Control")
if os.path.exists(f"{MOCK}/driver_home.png"):
    s.shapes.add_picture(f"{MOCK}/driver_home.png", Inches(0.9), Inches(1.55),
                         height=Inches(5.6))
if os.path.exists(f"{MOCK}/admin_dashboard.png"):
    s.shapes.add_picture(f"{MOCK}/admin_dashboard.png", Inches(4.2), Inches(1.55),
                         height=Inches(5.6))
bullets(s, Inches(7.7), Inches(1.8), Inches(5.2), Inches(5.3), [
    ("Driver: START TRIP dabao, done", "GPS apne aap share hota hai — screen off ho toh bhi"),
    ("Zero distraction while driving", "safety first — koi interaction ki zaroorat nahi"),
    ("Admin Dashboard", "buses, active trips, students, drivers — sab ek nazar mein"),
    ("Class-wise students", "Class 7-B ke sab bachche ek tap mein filter"),
    ("Trip History", "kaunsi bus kab chali, kitni der — poora record"),
    ("Activity logs", "kaun kya kab kiya — sab recorded, koi mita nahi sakta"),
], size=14, gap=11)
notes(s, "DRIVER+ADMIN (1.5 min): Driver safety par zor do — 'Driver ko driving ke time "
         "phone chhoona hi nahi padta.' Admin side: 'School ke paas full control hai.'")

# ═══ SLIDE 6 — Security (with single-device) ═══
s = slide()
title_bar(s, "SECURITY & PRIVACY", "Bachchon Ki Safety = #1 Priority 🔐")
rows = [
    ("Ek ID = Ek Phone (Single-Device Login)", "dusre phone par login karo toh pehla "
     "automatically logout — password sharing ka misuse impossible"),
    ("Admin ko dikhta hai kaun LOGGED IN hai", "har student/driver ka device status + "
     "Force Logout button"),
    ("Student sirf apni bus dekh sakta hai", "server-side rules — app hack karke bhi "
     "bypass nahi hota"),
    ("Koi public tracking link nahi", "bina school account ke kuch nahi dikhta"),
    ("Accounts sirf school banata hai", "koi khud register nahi kar sakta"),
    ("Data minimal & private", "location history store nahi hoti — trip khatam, data delete"),
    ("Passwords encrypted (Google Firebase)", "kabhi plain text mein store nahi hote"),
]
y = Inches(1.62)
for head, body in rows:
    box(s, Inches(0.6), y, Inches(12.1), Inches(0.72), fill=WHITE, line=DIVIDER)
    text(s, Inches(0.85), y + Inches(0.06), Inches(0.5), Inches(0.55),
         [[("🔒", 15, False, TEXT)]])
    text(s, Inches(1.4), y + Inches(0.04), Inches(4.9), Inches(0.6),
         [[(head, 13.5, True, TEXT)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(6.4), y + Inches(0.04), Inches(6.2), Inches(0.6),
         [[(body, 12, False, TEXT2)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.82)
notes(s, "SECURITY (2 min): Yeh slide parents/principal ke liye sabse important. "
         "Single-device feature highlight karo: 'WhatsApp jaisa — ek account ek hi phone "
         "par. Password share ho bhi jaye, ek time par ek hi device chalega, aur admin "
         "ko dikhega kaun kahan logged in hai.'")

# ═══ SLIDE 7 — Technology + Android/iOS ═══
s = slide()
title_bar(s, "TECHNOLOGY", "Modern Tech — Android + iPhone Dono Par")
tech = [
    ("Flutter", "Google ki technology — ek code se Android + iOS dono apps", PRIMARY),
    ("Firebase", "Google ka secure backend — login, database, real-time", ACCENT_DARK),
    ("OpenStreetMap", "Free world map — koi paid API nahi chahiye", SUCCESS),
    ("Real-time GPS", "Bus location har 5 second update hoti hai", PURPLE),
]
x = Inches(0.6)
for name, desc, col in tech:
    box(s, x, Inches(1.7), Inches(2.95), Inches(2.0), fill=WHITE, line=DIVIDER)
    tb = box(s, x, Inches(1.7), Inches(2.95), Inches(0.14), fill=col)
    text(s, x + Inches(0.22), Inches(2.0), Inches(2.5), Inches(0.5),
         [[(name, 18, True, TEXT)]])
    text(s, x + Inches(0.22), Inches(2.55), Inches(2.55), Inches(1.1),
         [[(desc, 12.5, False, TEXT2)]])
    x += Inches(3.13)
box(s, Inches(0.6), Inches(4.05), Inches(12.1), Inches(2.7), fill=PRIMARY_SOFT)
text(s, Inches(1.0), Inches(4.3), Inches(11.4), Inches(2.3), [
    [("Quality & Platform facts:", 16, True, PRIMARY_DARK)],
    [("✓  Android APK ready & tested — aaj hi install ho sakta hai", 14, False, TEXT)],
    [("✓  iOS (iPhone) — same code, sirf Apple Developer account chahiye "
      "($99/year, school ke naam par)", 14, False, TEXT)],
    [("✓  ~9,000 lines code · 0 analyzer errors · 15/15 automated tests passing", 14, False, TEXT)],
    [("✓  Running cost: chhote school ke liye approx ₹0/month (Google free tier)", 14, False, TEXT)],
    [("✓  Version 1.2.0 — single-device security, roll-number login, trip history", 14, False, TEXT)],
], space_after=7)
notes(s, "TECH (1 min): Simple rakho — 'Google ki hi technology par bana hai. Android "
         "phone ke liye APK ready hai. iPhone ke liye same code hai, bas Apple ka "
         "developer account chahiye jo school le sakta hai. Chalane ka kharcha lagbhag zero.'")

# ═══ SLIDE 8 — Live Demo ═══
s = slide()
title_bar(s, "LIVE DEMO", "Ab Main Chala Ke Dikhata Hoon 🎬")
steps = [
    ("1", "Admin login", "Dashboard — buses, students, active trips ek nazar mein"),
    ("2", "Student add karo", "Class dropdown + Roll number — 30 second ka kaam"),
    ("3", "Driver phone", "START TRIP dabaya — GPS sharing on"),
    ("4", "Student phone", "Class 7-B + Roll 23 + password se login → Track My Bus → LIVE! 🚌"),
    ("5", "Security demo", "Dusre phone par same ID login → pehla phone auto-logout!"),
    ("6", "Admin panel", "LOGGED IN status + Force Logout button dikhao"),
]
y = Inches(1.65)
for num, head, body in steps:
    box(s, Inches(1.2), y, Inches(10.9), Inches(0.8), fill=WHITE, line=DIVIDER)
    ovl = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.45), y + Inches(0.15),
                             Inches(0.5), Inches(0.5))
    ovl.fill.solid(); ovl.fill.fore_color.rgb = PRIMARY
    ovl.line.fill.background(); ovl.shadow.inherit = False
    tf = ovl.text_frame; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run(); r.text = num
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
    text(s, Inches(2.2), y + Inches(0.1), Inches(3.0), Inches(0.6),
         [[(head, 15.5, True, TEXT)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(5.3), y + Inches(0.1), Inches(6.5), Inches(0.6),
         [[(body, 13.5, False, TEXT2)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.9)
notes(s, "DEMO (3-4 min): 2 phones ready rakho. Step 5 ka single-device demo sabse "
         "impressive hai — audience ko live dikhao ki dusre phone par login karte hi "
         "pehla logout ho jaata hai. Backup video zaroor banao pehle se!")

# ═══ SLIDE 9 — Closing ═══
s = slide()
title_bar(s, "NEXT STEPS", "Launch Plan — Chhota Start, Bada Impact")
cols = [
    ("Phase 1 — Pilot (2 weeks)", PRIMARY,
     ["1 bus, 1 driver, ~20 students", "Feedback lo, improve karo"]),
    ("Phase 2 — Full Rollout", SUCCESS,
     ["Saari buses + routes", "Parents ko onboarding note"]),
    ("Phase 3 — Future", ACCENT_DARK,
     ["iPhone version launch", "Parent accounts", "Aur bhi features"]),
]
x = Inches(0.6)
for title_, col, items in cols:
    box(s, x, Inches(1.7), Inches(4.0), Inches(2.9), fill=WHITE, line=DIVIDER)
    hd = box(s, x, Inches(1.7), Inches(4.0), Inches(0.6), fill=col)
    text(s, x + Inches(0.2), Inches(1.78), Inches(3.6), Inches(0.45),
         [[(title_, 14.5, True, WHITE)]])
    ty = Inches(2.5)
    for it in items:
        text(s, x + Inches(0.3), ty, Inches(3.5), Inches(0.6),
             [[("•  " + it, 13.5, False, TEXT)]])
        ty += Inches(0.55)
    x += Inches(4.25)
big = box(s, Inches(0.6), Inches(5.0), Inches(12.1), Inches(1.7), fill=PRIMARY)
text(s, Inches(0.6), Inches(5.3), Inches(12.1), Inches(1.3), [
    [("Never Miss Bus — kyunki har bachche ka time, safety aur parents ka "
      "sukoon important hai.", 19, True, WHITE)],
    [("Questions?  Demo ready hai! 🚌  |  Android + iOS  |  Secure  |  ~₹0/month",
      15, False, RGBColor(0xCF, 0xE0, 0xFF))],
], align=PP_ALIGN.CENTER, space_after=10)
notes(s, "CLOSING (30 sec): Confident end karo. Pilot proposal do — 'Ek bus se shuru "
         "karte hain, 2 hafte mein results dikha denge.' Phir Q&A. All the best! 🚀")

prs.save(OUT)
print("saved:", OUT, os.path.getsize(OUT)//1024, "KB")
